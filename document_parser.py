"""Parse PDF, PPTX, and DOCX files into plain text."""

from __future__ import annotations

import io
from pathlib import Path


def parse_pdf(file_bytes: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(file_bytes))
    pages = [page.extract_text() or "" for page in reader.pages]
    return "\n\n".join(pages)


def parse_pptx(file_bytes: bytes) -> str:
    from pptx import Presentation

    prs = Presentation(io.BytesIO(file_bytes))
    slides = []
    for i, slide in enumerate(prs.slides, 1):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    line = "".join(run.text for run in para.runs).strip()
                    if line:
                        texts.append(line)
        if texts:
            slides.append(f"[Slide {i}]\n" + "\n".join(texts))
    return "\n\n".join(slides)


def parse_docx(file_bytes: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(file_bytes))
    paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
    return "\n\n".join(paragraphs)


def parse_uploaded_file(name: str, file_bytes: bytes) -> str:
    suffix = Path(name).suffix.lower()
    if suffix == ".pdf":
        return parse_pdf(file_bytes)
    elif suffix == ".pptx":
        return parse_pptx(file_bytes)
    elif suffix == ".docx":
        return parse_docx(file_bytes)
    else:
        raise ValueError(f"Unsupported file type: {suffix}")
